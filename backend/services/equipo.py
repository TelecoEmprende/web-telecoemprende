import base64
from io import BytesIO

import psycopg2
from fpdf import FPDF
from flask import session
from openpyxl import Workbook
from openpyxl.drawing.image import Image as XLImage
from werkzeug.security import check_password_hash, generate_password_hash

from backend.services.db import ids_de_personas
from backend.config import (
    CARGOS_VALIDOS,
    DATABASE_URL,
    EQUIPOS_VALIDOS,
)
from backend.services.registrations import _celda_segura

# Hash "de relleno" para cuando el email no existe: sin esto, saltarse
# check_password_hash en ese caso haría que la respuesta fuera más rápida
# para emails no registrados, filtrando por tiempo qué correos están de alta.
_DUMMY_HASH = generate_password_hash("no-existe-ningun-usuario-con-este-email")


def _get_connection():
    return psycopg2.connect(DATABASE_URL)


# La inscripción de la que sale una cuenta nueva: la más reciente con su email.
# Quien se inscribió con otro correo se vincula a mano (`registro_id` en
# `actualizar_equipo_acceso`).
_INSCRIPCION = (
    "(SELECT id FROM registrations WHERE lower(email) = %s ORDER BY created_at DESC LIMIT 1)"
)


def listar_directorio_club() -> list[dict]:
    """Quién es quién del club entero, para el widget de "Mi semana" -- solo
    lo básico (nombre, equipos, cargo). Nada de notas ni datos personales
    (ver `listar_equipo_accesos`, la versión completa de admin).
    """
    with _get_connection() as conn:
        with conn.cursor() as cur:
            cur.execute(
                """
                SELECT email, equipos, vp_de, cargo, nombre
                FROM equipo_accesos WHERE activo = TRUE ORDER BY nombre, email
                """
            )
            filas = cur.fetchall()

    return [
        {
            "email": f[0], "equipos": f[1], "vp_de": f[2], "cargo": f[3],
            "nombre": f[4],
        }
        for f in filas
    ]


def ordenar_equipos(equipos: list[str]) -> list[str]:
    """Sin duplicados y en el orden en que llegan: el primero es la 1ª
    preferencia de la persona y el segundo la 2ª (no se ordena alfabéticamente,
    que perdería esa información)."""
    return list(dict.fromkeys(equipos))


def login_equipo(email: str, password: str) -> dict | None:
    email = email.strip().lower()
    with _get_connection() as conn:
        with conn.cursor() as cur:
            cur.execute(
                """
                SELECT password_hash, equipos, vp_de, cargo, nombre,
                       COALESCE(es_admin, FALSE)
                FROM equipo_accesos WHERE email = %s AND activo = TRUE
                """,
                (email,),
            )
            row = cur.fetchone()

    password_hash = row[0] if row is not None else _DUMMY_HASH
    password_ok = check_password_hash(password_hash, password)

    if row is None or not password_ok:
        return None

    equipos = [e for e in row[1] if e in EQUIPOS_VALIDOS]
    vp_de = [e for e in row[2] if e in equipos]
    cargo = row[3] if row[3] in CARGOS_VALIDOS else ""
    nombre = row[4]
    es_admin = bool(row[5])

    # session.clear() por higiene ante fijación de sesión.
    session.clear()
    session.permanent = True
    session["equipo_email"] = email
    session["equipo_teams"] = equipos
    session["equipo_vp_de"] = vp_de
    session["equipo_cargo"] = cargo
    session["equipo_nombre"] = nombre
    session["equipo_admin"] = es_admin
    # Solo quien tiene el permiso de admin marcado (no el departamento ni el
    # cargo) recibe la sesión de admin: grupo Admin y rutas /api/admin/*.
    if es_admin:
        session["admin_auth"] = True
    return {
        "teams": equipos, "vp_de": vp_de, "cargo": cargo, "nombre": nombre,
        "admin": es_admin,
    }


def is_equipo_authenticated() -> bool:
    return "equipo_email" in session


def equipo_session_info() -> dict:
    return {
        "teams": session.get("equipo_teams", []),
        "vp_de": session.get("equipo_vp_de", []),
        "cargo": session.get("equipo_cargo", ""),
        "email": session.get("equipo_email", ""),
        "nombre": session.get("equipo_nombre", ""),
        "admin": session.get("equipo_admin", False),
    }


def logout_equipo() -> None:
    session.clear()



def listar_equipo_accesos() -> list[dict]:
    with _get_connection() as conn:
        with conn.cursor() as cur:
            cur.execute(
                """
                SELECT id, email, equipos, vp_de, cargo, activo, created_at,
                       tags, notas, nombre, dni, correo_personal,
                       foto, apellidos, COALESCE(es_admin, FALSE),
                       COALESCE(en_web, FALSE), registro_id
                FROM equipo_accesos ORDER BY email
                """
            )
            filas = cur.fetchall()

    return [
        {
            "id": f[0],
            "email": f[1],
            "equipos": f[2],
            "vp_de": f[3],
            "cargo": f[4],
            "activo": f[5],
            "created_at": f[6].isoformat(),
            "tags": f[7],
            "notas": f[8],
            "nombre": f[9],
            "dni": f[10],
            "correo_personal": f[11],
            "foto": f[12],
            "apellidos": f[13],
            "es_admin": f[14],
            "en_web": f[15],
            "registro_id": f[16],
        }
        for f in filas
    ]


# Orden en la web: presidencia, board, VPs y luego el resto.
_RANGO_WEB = {"presidente": 0, "boardmember": 1, "vicepresidente": 2}

_DEPTO_WEB = {
    "marketing": ("Marketing", "Marketing"),
    "eventos": ("Eventos", "Events"),
    "ingenieria": ("Ingeniería", "Engineering"),
}


def puesto_web(cargo: str, vp_de: list[str]) -> dict | None:
    """El puesto que sale bajo el nombre en la web, en español e inglés: solo
    para presidencia, board y VPs (de su departamento); el resto, nada."""
    if cargo == "presidente":
        return {"es": "Presidente", "en": "President"}
    if cargo == "boardmember":
        return {"es": "Board member", "en": "Board member"}
    deptos = [d for d in vp_de if d in _DEPTO_WEB]
    if deptos:
        return {
            "es": "VP de " + " y ".join(_DEPTO_WEB[d][0] for d in deptos),
            "en": "VP of " + " & ".join(_DEPTO_WEB[d][1] for d in deptos),
        }
    if cargo == "vicepresidente":
        return {"es": "Vicepresidente", "en": "Vice President"}
    return None


def listar_miembros_web() -> list[dict]:
    """Quién sale en la sección «Equipo» de la web pública: cuentas marcadas
    `en_web` y con foto subida. Ruta pública, así que solo viaja nombre,
    primer apellido, foto y, si lo tiene, su puesto (`puesto_web`); nada de
    email, DNI ni departamentos.
    """
    with _get_connection() as conn:
        with conn.cursor() as cur:
            cur.execute(
                """
                SELECT nombre, apellidos, foto, cargo, vp_de
                FROM equipo_accesos
                WHERE en_web AND nombre <> '' AND foto <> ''
                """
            )
            filas = cur.fetchall()

    filas.sort(key=lambda f: (_RANGO_WEB.get(f[3], 2 if f[4] else 3), f[0].lower()))
    return [
        {
            "nombre": f[0],
            "apellido": (f[1].split() or [""])[0],
            "foto": f[2],
            "puesto": puesto_web(f[3], f[4]),
        }
        for f in filas
    ]


def generar_pdf_equipo_en_memoria() -> BytesIO:
    """Listado de todos los accesos de equipo en PDF, para imprimir o
    archivar (p. ej. registro de socios) -- ver `api/admin.py`."""
    accesos = listar_equipo_accesos()

    pdf = FPDF(orientation="L", unit="mm", format="A4")
    pdf.set_auto_page_break(auto=True, margin=12)
    pdf.add_page()
    pdf.set_font("Helvetica", "B", 14)
    pdf.cell(0, 10, "TelecoEmprende -- Miembros del equipo", new_x="LMARGIN", new_y="NEXT")

    columnas = [
        ("Nombre", 40), ("Email", 55), ("DNI", 25), ("Correo personal", 55),
        ("Equipos", 45), ("Cargo", 25), ("Activo", 15),
    ]
    pdf.set_font("Helvetica", "B", 9)
    for titulo, ancho in columnas:
        pdf.cell(ancho, 8, titulo, border=1)
    pdf.ln()

    pdf.set_font("Helvetica", "", 9)
    for acceso in accesos:
        valores = [
            acceso["nombre"] or "-",
            acceso["email"],
            acceso["dni"] or "-",
            acceso["correo_personal"] or "-",
            ", ".join(acceso["equipos"]) or "-",
            acceso["cargo"] or "-",
            "Sí" if acceso["activo"] else "No",
        ]
        for (_, ancho), valor in zip(columnas, valores):
            texto = valor.encode("latin-1", "replace").decode("latin-1")
            pdf.cell(ancho, 7, texto, border=1)
        pdf.ln()

    return BytesIO(bytes(pdf.output()))


def generar_excel_equipo_en_memoria() -> BytesIO:
    """Datos personales de todo el equipo (lo que cada cual rellena en
    `/equipo/datosformulario`) con sus departamentos y cargo, y la foto
    incrustada en la última columna -- ver `api/admin.py`."""
    accesos = listar_equipo_accesos()

    wb = Workbook()
    ws = wb.active
    ws.title = "Equipo"
    ws.append([
        "Nombre", "Apellidos", "DNI/NIE/Pasaporte", "Correo de contacto",
        "Email de acceso", "Departamentos", "VP de", "Cargo", "Activo", "Foto",
    ])
    for col, ancho in zip("ABCDEFGHIJ", (20, 25, 18, 32, 32, 30, 20, 14, 8, 12)):
        ws.column_dimensions[col].width = ancho

    for fila, a in enumerate(accesos, start=2):
        ws.append([
            _celda_segura(a["nombre"]),
            _celda_segura(a["apellidos"]),
            _celda_segura(a["dni"]),
            _celda_segura(a["correo_personal"]),
            _celda_segura(a["email"]),
            ", ".join(a["equipos"]),
            ", ".join(a["vp_de"]),
            a["cargo"],
            "Sí" if a["activo"] else "No",
        ])
        if a["foto"]:
            # La foto ya pasó `_foto` al guardarse (data URL validado), así que
            # aquí basta con decodificarla.
            imagen = XLImage(BytesIO(base64.b64decode(a["foto"].split(",", 1)[1])))
            imagen.width = imagen.height = 64
            ws.add_image(imagen, f"J{fila}")
            ws.row_dimensions[fila].height = 50

    output = BytesIO()
    wb.save(output)
    output.seek(0)
    return output


def datos_formulario(email: str) -> dict | None:
    with _get_connection() as conn:
        with conn.cursor() as cur:
            cur.execute(
                """
                SELECT nombre, apellidos, dni, correo_personal, foto
                FROM equipo_accesos WHERE email = %s
                """,
                (email.strip().lower(),),
            )
            f = cur.fetchone()
    if f is None:
        return None
    return {"nombre": f[0], "apellidos": f[1], "dni": f[2], "correo": f[3], "foto": f[4]}


def guardar_datos_formulario(
    email: str, nombre: str, apellidos: str, dni: str, correo: str, foto: str | None
) -> bool:
    """Ficha que rellena la propia persona. `foto=None` deja la que ya tenía
    (no hace falta volver a subirla para corregir un apellido)."""
    campos = ["nombre = %s", "apellidos = %s", "dni = %s", "correo_personal = %s"]
    valores: list = [nombre, apellidos, dni, correo]
    if foto is not None:
        campos.append("foto = %s")
        valores.append(foto)
    valores.append(email.strip().lower())
    with _get_connection() as conn:
        with conn.cursor() as cur:
            cur.execute(
                f"UPDATE equipo_accesos SET {', '.join(campos)} WHERE email = %s",
                valores,
            )
            actualizado = cur.rowcount > 0
        conn.commit()
    return actualizado


def miembros_activos(equipo: str) -> list[dict]:
    """Los accesos activos de un departamento -- filtro que reutilizan el
    directorio de marketing (`api_miembros`) y `salud_equipo`, para no tener
    la misma condición escrita dos veces."""
    return [a for a in listar_equipo_accesos() if equipo in a["equipos"] and a["activo"]]


def _acceso_valido(equipos: list[str], vp_de: list[str], cargo: str) -> bool:
    """Un acceso necesita al menos un departamento O un cargo de dirección.

    El cargo a solas (board member sin departamento) es válido a propósito:
    el Inicio, las Notas y el club se ven igual. Sin ninguno de los dos, en
    cambio, la cuenta no daría acceso a nada. El mínimo de dos departamentos
    se exige en la API al cambiarlos, no aquí: así una fila antigua con uno
    solo se puede seguir editando (cargo, VP...) sin tener que arreglarla
    primero.
    """
    if cargo and cargo not in CARGOS_VALIDOS:
        return False
    if any(e not in EQUIPOS_VALIDOS for e in equipos):
        return False
    if not equipos and not cargo:
        return False
    return all(v in equipos for v in vp_de)


def crear_equipo_acceso(
    email: str,
    password: str,
    equipos: list[str],
    vp_de: list[str] | None = None,
    cargo: str = "",
    nombre: str = "",
    es_admin: bool = False,
    apellidos: str = "",
) -> dict | None:
    """Devuelve None si el email ya existe o si equipos/vp_de/cargo no son válidos."""
    email = email.strip().lower()
    equipos = ordenar_equipos(equipos)
    vp_de = sorted(set(vp_de or []))
    cargo = cargo or ""
    nombre = (nombre or "").strip()

    if not _acceso_valido(equipos, vp_de, cargo):
        return None

    with _get_connection() as conn:
        with conn.cursor() as cur:
            cur.execute("SELECT 1 FROM equipo_accesos WHERE email = %s", (email,))
            if cur.fetchone() is not None:
                return None

            cur.execute(
                f"""
                INSERT INTO equipo_accesos
                    (email, password_hash, equipos, vp_de, cargo, nombre, es_admin, apellidos,
                     registro_id)
                VALUES (%s, %s, %s, %s, %s, %s, %s, %s, {_INSCRIPCION})
                RETURNING id, email, equipos, vp_de, cargo, activo, created_at, nombre,
                          es_admin, apellidos
                """,
                (
                    email, generate_password_hash(password), equipos, vp_de, cargo, nombre,
                    es_admin, (apellidos or "").strip(), email,
                ),
            )
            fila = cur.fetchone()
        conn.commit()

    return {
        "id": fila[0],
        "email": fila[1],
        "equipos": fila[2],
        "vp_de": fila[3],
        "cargo": fila[4],
        "activo": fila[5],
        "created_at": fila[6].isoformat(),
        "nombre": fila[7],
        "es_admin": fila[8],
        "apellidos": fila[9],
    }


def registrar_equipo_acceso(email: str, password: str) -> bool:
    """Alta que se pide la propia persona: cuenta desactivada y sin equipos.

    No pasa por `_acceso_valido` a propósito -- la cuenta nace justamente sin
    nada (equipos '{}', cargo ''), y con `activo = FALSE` no puede iniciar
    sesión (ver `login_equipo`), así que no da acceso a nada hasta que admin le
    asigna equipos/cargo y la activa desde el panel.

    `ON CONFLICT DO NOTHING` en vez de SELECT + INSERT: una sola sentencia, sin
    carrera entre dos altas del mismo email a la vez. Devuelve False si ya
    existía.

    ponytail: alta temporal mientras entra el equipo. Para quitarla, borrar
    esta función, la ruta POST /api/equipo/registro y el modo "crear cuenta"
    del formulario de login.
    """
    with _get_connection() as conn:
        with conn.cursor() as cur:
            cur.execute(
                f"""
                INSERT INTO equipo_accesos (email, password_hash, activo, registro_id)
                VALUES (%s, %s, FALSE, {_INSCRIPCION})
                ON CONFLICT (email) DO NOTHING
                """,
                (email.strip().lower(), generate_password_hash(password), email.strip().lower()),
            )
            creado = cur.rowcount > 0
        conn.commit()

    return creado


def actualizar_equipo_acceso(
    acceso_id: int,
    equipos: list[str] | None = None,
    vp_de: list[str] | None = None,
    cargo: str | None = None,
    activo: bool | None = None,
    password: str | None = None,
    nombre: str | None = None,
    apellidos: str | None = None,
    dni: str | None = None,
    correo_personal: str | None = None,
    email: str | None = None,
    registro_id: int | None = None,
    es_admin: bool | None = None,
    foto: str | None = None,
    en_web: bool | None = None,
) -> bool:
    """Actualiza solo los campos que se pasan. Devuelve False si el id no existe
    o si equipos/vp_de/cargo no son válidos.

    `vp_de` se valida contra `equipos`: si se cambia uno sin el otro en la misma
    llamada, se valida contra el `equipos` ya guardado en la fila.
    """
    if equipos is not None:
        equipos = ordenar_equipos(equipos)
    if vp_de is not None:
        vp_de = sorted(set(vp_de))
    campos = []
    valores: list = []

    # equipos, vp_de y cargo se validan juntos: cambiar uno solo puede dejar el
    # acceso inválido (quitarle el cargo a un board member sin departamento lo
    # dejaría sin acceso a nada), así que se comprueba el trío ya resuelto
    # contra lo que hay guardado.
    if equipos is not None or vp_de is not None or cargo is not None:
        with _get_connection() as conn:
            with conn.cursor() as cur:
                cur.execute(
                    "SELECT equipos, vp_de, cargo FROM equipo_accesos WHERE id = %s",
                    (acceso_id,),
                )
                fila = cur.fetchone()
        if fila is None:
            return False

        if not _acceso_valido(
            equipos if equipos is not None else fila[0],
            vp_de if vp_de is not None else fila[1],
            cargo if cargo is not None else fila[2],
        ):
            return False

        if equipos is not None:
            campos.append("equipos = %s")
            valores.append(equipos)
        if vp_de is not None:
            campos.append("vp_de = %s")
            valores.append(vp_de)
        if cargo is not None:
            campos.append("cargo = %s")
            valores.append(cargo)
    if activo is not None:
        campos.append("activo = %s")
        valores.append(activo)
    if password:
        campos.append("password_hash = %s")
        valores.append(generate_password_hash(password))
    if nombre is not None:
        campos.append("nombre = %s")
        valores.append(nombre.strip())
    if apellidos is not None:
        campos.append("apellidos = %s")
        valores.append(apellidos.strip())
    if dni is not None:
        campos.append("dni = %s")
        valores.append(dni.strip())
    if correo_personal is not None:
        campos.append("correo_personal = %s")
        valores.append(correo_personal.strip().lower())
    if email is not None:
        # El email es el login: único (la API traduce el choque a un 400).
        campos.append("email = %s")
        valores.append(email.strip().lower())
    if registro_id is not None:
        # 0 desvincula la cuenta de su inscripción.
        campos.append("registro_id = %s")
        valores.append(registro_id or None)
    if es_admin is not None:
        campos.append("es_admin = %s")
        valores.append(es_admin)
    if foto is not None:
        # Ya validada con `_foto` en la API; "" vuelve a la de `public/`.
        campos.append("foto = %s")
        valores.append(foto)
    if en_web is not None:
        campos.append("en_web = %s")
        valores.append(en_web)

    if not campos:
        return False

    valores.append(acceso_id)
    with _get_connection() as conn:
        with conn.cursor() as cur:
            cur.execute(
                f"UPDATE equipo_accesos SET {', '.join(campos)} WHERE id = %s",
                valores,
            )
            actualizado = cur.rowcount > 0
        conn.commit()

    return actualizado


def eliminar_equipo_acceso(acceso_id: int) -> bool:
    with _get_connection() as conn:
        with conn.cursor() as cur:
            cur.execute("DELETE FROM equipo_accesos WHERE id = %s", (acceso_id,))
            eliminado = cur.rowcount > 0
        conn.commit()

    return eliminado


# ---------------------------------------------------------------------------
# Calendario compartido: lo gestiona admin, lo ve cualquiera logueado en /equipo.
# ---------------------------------------------------------------------------


def listar_eventos_calendario() -> list[dict]:
    with _get_connection() as conn:
        with conn.cursor() as cur:
            cur.execute(
                """
                SELECT id, titulo, descripcion, fecha, hora, confirmados, asistio
                FROM calendario_eventos_v ORDER BY fecha, hora
                """
            )
            filas = cur.fetchall()

    return [
        {
            "id": f[0],
            "titulo": f[1],
            "descripcion": f[2],
            "fecha": f[3].isoformat(),
            "hora": f[4],
            "confirmados": f[5],
            "asistio": f[6],
        }
        for f in filas
    ]


def confirmar_evento_calendario(evento_id: int, email: str, confirmar: bool) -> bool:
    """"Voy" / "no voy" de la propia persona a un evento del club."""
    return _marcar_en_evento("evento_confirmados", evento_id, email, confirmar)


def marcar_asistio_evento(evento_id: int, email: str, asistio: bool) -> bool:
    """Check-in real el día del evento -- lo marca quien gestiona la puerta
    (VP/admin, ver `_puede_editar_calendario_club`), no la propia persona.
    Alimenta `metricas_club` ("asistencia media"), a diferencia de
    `evento_confirmados`, que es solo la intención previa."""
    return _marcar_en_evento("evento_asistentes", evento_id, email, asistio)


def _marcar_en_evento(tabla: str, evento_id: int, email: str, poner: bool) -> bool:
    """Pone o quita a una persona de una lista de un evento. False si el
    evento no existe."""
    with _get_connection() as conn:
        with conn.cursor() as cur:
            cur.execute("SELECT 1 FROM calendario_eventos WHERE id = %s", (evento_id,))
            if cur.fetchone() is None:
                return False
            (persona_id,) = ids_de_personas(cur, [email])
            if poner:
                cur.execute(
                    f"INSERT INTO {tabla} (evento_id, persona_id) VALUES (%s, %s)"
                    " ON CONFLICT DO NOTHING",
                    (evento_id, persona_id),
                )
            else:
                cur.execute(
                    f"DELETE FROM {tabla} WHERE evento_id = %s AND persona_id = %s",
                    (evento_id, persona_id),
                )
        conn.commit()
    return True


def crear_evento_calendario(titulo: str, descripcion: str, fecha: str, hora: str) -> dict:
    with _get_connection() as conn:
        with conn.cursor() as cur:
            cur.execute(
                """
                INSERT INTO calendario_eventos (titulo, descripcion, fecha, hora)
                VALUES (%s, %s, %s, %s)
                RETURNING id, titulo, descripcion, fecha, hora
                """,
                (titulo.strip(), descripcion.strip(), fecha, hora.strip()),
            )
            fila = cur.fetchone()
        conn.commit()

    return {
        "id": fila[0],
        "titulo": fila[1],
        "descripcion": fila[2],
        "fecha": fila[3].isoformat(),
        "hora": fila[4],
    }


def actualizar_evento_calendario(
    evento_id: int, titulo: str, descripcion: str, fecha: str, hora: str
) -> bool:
    with _get_connection() as conn:
        with conn.cursor() as cur:
            cur.execute(
                """
                UPDATE calendario_eventos
                SET titulo = %s, descripcion = %s, fecha = %s, hora = %s
                WHERE id = %s
                """,
                (titulo.strip(), descripcion.strip(), fecha, hora.strip(), evento_id),
            )
            actualizado = cur.rowcount > 0
        conn.commit()

    return actualizado


def eliminar_evento_calendario(evento_id: int) -> bool:
    with _get_connection() as conn:
        with conn.cursor() as cur:
            cur.execute("DELETE FROM calendario_eventos WHERE id = %s", (evento_id,))
            eliminado = cur.rowcount > 0
        conn.commit()

    return eliminado


def actualizar_perfil(
    email: str,
    tags: list[str] | None = None,
    notas: str | None = None,
    foto: str | None = None,
) -> bool:
    """Etiquetas de habilidad, nota y foto de una persona.

    Separado de `actualizar_equipo_acceso` a propósito: eso son permisos y solo
    lo toca quien administra; esto es contexto de trabajo y lo edita cualquiera
    del departamento desde su ficha. Mezclarlos habría significado exponer los
    permisos a quien solo quiere apuntar "entrega rápido pero mejor una cosa a
    la vez".
    """
    campos, valores = [], []
    if tags is not None:
        campos.append("tags = %s")
        valores.append(sorted({t.strip() for t in tags if t.strip()}))
    if notas is not None:
        campos.append("notas = %s")
        valores.append(notas.strip())
    if foto is not None:
        # "" borra la foto propia y devuelve a la de `public/`, que es la
        # única forma de deshacer una subida.
        campos.append("foto = %s")
        valores.append(foto)

    if not campos:
        return False

    valores.append(email.strip().lower())
    with _get_connection() as conn:
        with conn.cursor() as cur:
            cur.execute(
                f"UPDATE equipo_accesos SET {', '.join(campos)} WHERE email = %s",
                valores,
            )
            actualizado = cur.rowcount > 0
        conn.commit()

    return actualizado
