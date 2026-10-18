"""Perfil propio (/api/equipo/perfil), cambio de contraseña y el Excel de /admin."""

import base64
import os
import unittest
from io import BytesIO

os.environ["DATABASE_URL"] = os.environ.get(
    "TEST_DATABASE_URL",
    "postgresql://telecoemprende:telecoemprende@localhost:5432/telecoemprende_test",
)

from openpyxl import load_workbook  # noqa: E402

import app  # noqa: E402
import backend.services.equipo as equipo_service  # noqa: E402
import backend.services.security as security_service  # noqa: E402
from backend.services.db import migrar  # noqa: E402

# PNG de 1x1 válido.
FOTO = (
    "data:image/png;base64,iVBORw0KGgoAAAANSUhEUgAAAAEAAAABCAYAAAAfFcSJAAAADUlEQVR4"
    "nGNgYGD4DwABBAEAwS2OUAAAAABJRU5ErkJggg=="
)


class DatosFormularioTestCase(unittest.TestCase):
    def setUp(self):
        security_service.request_log.clear()
        migrar()

        conn = equipo_service._get_connection()
        with conn.cursor() as cur:
            cur.execute("DELETE FROM members")
        conn.commit()
        conn.close()

        equipo_service.crear_equipo_acceso(
            "ana@alumnos.upm.es", "contrasena-larga", ["marketing"], vp_de=["marketing"]
        )
        app.app.config["TESTING"] = True
        self.client = app.app.test_client()

    def entrar(self):
        self.client.post(
            "/api/equipo/login",
            json={"email": "ana@alumnos.upm.es", "password": "contrasena-larga"},
        )

    def guardar(self, **cambios):
        datos = {
            "nombre": "Ana", "apellidos": "García López", "dni": "12345678-z",
            "linkedin": "https://www.linkedin.com/in/ana", "foto": FOTO, **cambios,
        }
        return self.client.put("/api/equipo/perfil", json=datos)

    def test_sin_sesion_no_deja_leer_ni_guardar(self):
        self.assertEqual(self.client.get("/api/equipo/perfil").status_code, 401)
        self.assertEqual(self.guardar().status_code, 401)
        self.assertEqual(
            self.client.put("/api/equipo/password", json={"actual": "x", "nueva": "y" * 8}).status_code,
            401,
        )

    def test_guarda_y_devuelve_los_datos_propios(self):
        self.entrar()
        self.assertEqual(self.guardar().status_code, 200)

        datos = self.client.get("/api/equipo/perfil").get_json()
        self.assertEqual(datos["apellidos"], "García López")
        self.assertEqual(datos["dni"], "12345678Z")
        self.assertEqual(datos["linkedin"], "https://www.linkedin.com/in/ana")
        self.assertEqual(self.client.get(datos["foto"]).data, base64.b64decode(FOTO.split(",")[1]))

        # Sin `foto` se conserva la guardada.
        self.client.put("/api/equipo/perfil", json={
            "nombre": "Ana", "apellidos": "García", "dni": "X1234567L",
        })
        self.assertEqual(self.client.get("/api/equipo/perfil").get_json()["foto"], datos["foto"])
        # El nombre nuevo se ve ya en la sesión, sin volver a entrar.
        self.assertEqual(self.client.get("/api/equipo/session").get_json()["nombre"], "Ana")

    def test_el_perfil_no_toca_cargo_ni_departamentos_ni_permisos(self):
        self.entrar()
        self.guardar(cargo="presidente", equipos=["ingenieria"], es_admin=True, en_web=True)

        yo = equipo_service.listar_equipo_accesos()[0]
        self.assertEqual(yo["cargo"], "")
        self.assertEqual(yo["equipos"], ["marketing", "eventos", "ingenieria"])
        self.assertFalse(yo["es_admin"])
        self.assertFalse(yo["en_web"])

    def test_el_dni_no_sale_en_el_directorio_ni_en_la_ficha(self):
        self.entrar()
        self.guardar()
        equipo_service.crear_equipo_acceso("otra@alumnos.upm.es", "contrasena-larga", ["eventos"])

        ficha = self.client.get("/api/equipo/miembros/ficha?email=ana@alumnos.upm.es").get_json()
        self.assertNotIn("dni", ficha["ficha"])
        lista = self.client.get("/api/equipo/miembros").get_json()["miembros"]
        self.assertTrue(all("dni" not in m for m in lista))
        # Ni siquiera el de otra persona por el perfil: siempre es el propio.
        self.assertEqual(self.client.get("/api/equipo/perfil").get_json()["email"], "ana@alumnos.upm.es")

    def test_rechaza_datos_invalidos(self):
        self.entrar()
        self.assertEqual(self.guardar(apellidos="").status_code, 400)
        self.assertEqual(self.guardar(dni="12<3>").status_code, 400)
        self.assertEqual(self.guardar(linkedin="javascript:alert(1)").status_code, 400)
        self.assertEqual(self.guardar(foto="data:text/html;base64,PHNjcmlwdD4=").status_code, 400)
        # El documento es opcional: vacío se acepta.
        self.assertEqual(self.guardar(dni="").status_code, 200)

    def test_cambia_su_contrasena_pidiendo_la_actual(self):
        self.entrar()
        mal = self.client.put("/api/equipo/password", json={"actual": "incorrecta", "nueva": "otra-contrasena"})
        self.assertEqual(mal.status_code, 400)
        corta = self.client.put("/api/equipo/password", json={"actual": "contrasena-larga", "nueva": "corta"})
        self.assertEqual(corta.status_code, 400)

        bien = self.client.put(
            "/api/equipo/password", json={"actual": "contrasena-larga", "nueva": "otra-contrasena"}
        )
        self.assertEqual(bien.status_code, 200)

        self.client.post("/api/equipo/logout")
        vieja = self.client.post(
            "/api/equipo/login", json={"email": "ana@alumnos.upm.es", "password": "contrasena-larga"}
        )
        self.assertEqual(vieja.status_code, 401)
        nueva = self.client.post(
            "/api/equipo/login", json={"email": "ana@alumnos.upm.es", "password": "otra-contrasena"}
        )
        self.assertEqual(nueva.status_code, 200)

    def test_sin_sesion_de_admin_fetch_recibe_json_y_el_navegador_una_pagina(self):
        # `fetch` manda Accept: */* y `apiRequest` solo entiende JSON.
        fetch = self.client.get("/api/admin/equipo", headers={"Accept": "*/*"})
        self.assertEqual((fetch.status_code, fetch.get_json()["ok"]), (401, False))
        navegador = self.client.get(
            "/api/admin/equipo/excel", headers={"Accept": "text/html,application/xhtml+xml,*/*;q=0.8"}
        )
        self.assertEqual(navegador.status_code, 401)
        self.assertIn("text/html", navegador.content_type)

    def test_excel_solo_para_admin_con_roles_y_foto(self):
        self.assertEqual(self.client.get("/api/admin/equipo/excel").status_code, 401)

        self.entrar()
        self.guardar(nombre="=HYPERLINK()")
        with self.client.session_transaction() as s:
            s["admin_auth"] = True
        respuesta = self.client.get("/api/admin/equipo/excel")
        self.assertEqual(respuesta.status_code, 200)

        ws = load_workbook(BytesIO(respuesta.data)).active
        fila = [c.value for c in ws[2]]
        self.assertEqual(fila[0], "'=HYPERLINK()")
        self.assertEqual(fila[1:5], ["García López", "12345678Z", "ana@alumnos.upm.es",
                                     "marketing, eventos, ingenieria"])
        self.assertEqual(fila[5], "marketing")
        self.assertEqual(len(ws._images), 1)


if __name__ == "__main__":
    unittest.main()
